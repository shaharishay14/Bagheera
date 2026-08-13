"""Build 'Detailed Design - Bagheera (rev. 2, as built)'.

Mirrors the formatting of the original submitted document: 24 pt bold title,
13.5 pt bold section headings, default body text with bold lead-ins, bordered
tables, and figures.
"""
from __future__ import annotations

import os
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = "/Users/shaharishay/Projects/Bagheera"
OUT = os.path.join(REPO, "docs", "Detailed Design - Bagheera (rev 2, as built).docx")

doc = Document()

# Page setup matching the original (Letter, 1.25" side margins, 1" top)
sec = doc.sections[0]
sec.left_margin = Inches(1.25)
sec.right_margin = Inches(1.25)
sec.top_margin = Inches(1.0)
sec.bottom_margin = Inches(1.0)

CODE_GREY = RGBColor(0x44, 0x44, 0x44)


# --------------------------------------------------------------------------
# helpers
# --------------------------------------------------------------------------
def para(*parts, space_after=6, space_before=0, align=None, style=None):
    """parts: str  OR  (text, 'b'|'i'|'bi'|'c'|'u') tuples."""
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if align:
        p.alignment = align
    for part in parts:
        if isinstance(part, tuple):
            text, fmt = part
        else:
            text, fmt = part, ""
        r = p.add_run(text)
        if "b" in fmt:
            r.bold = True
        if "i" in fmt:
            r.italic = True
        if "u" in fmt:
            r.underline = True
        if "c" in fmt:
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
            r.font.color.rgb = CODE_GREY
    return p


def title(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(24)
    return p


def h(text, *, before=14, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(13.5)
    return p


def bullet(lead, rest="", *, code=False):
    """A bulleted item with an optional bold lead-in."""
    p = doc.add_paragraph(style="List Paragraph")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Inches(0.35)
    r = p.add_run("• ")
    if lead:
        r = p.add_run(lead)
        r.bold = True
    if rest:
        r = p.add_run(rest)
        if code:
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
    return p


def numbered(text):
    p = doc.add_paragraph(text, style="List Paragraph")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Inches(0.35)
    return p


def code(lines, size=8.5):
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.left_indent = Inches(0.25)
        r = p.add_run(line if line else " ")
        r.font.name = "Consolas"
        r.font.size = Pt(size)
        r.font.color.rgb = CODE_GREY
    doc.paragraphs[-1].paragraph_format.space_after = Pt(8)


def table(rows, *, widths=None, size=9, header=True):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.autofit = True
    for i, row in enumerate(rows):
        for j, cell_text in enumerate(row):
            cell = t.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            r = p.add_run(str(cell_text))
            r.font.size = Pt(size)
            if header and i == 0:
                r.bold = True
    if widths:
        for j, w in enumerate(widths):
            for row in t.rows:
                row.cells[j].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


FIGURES_DIR = os.path.join(REPO, "docs", "figures")


def figure(png, caption, width=6.4, *, placeholder=None):
    """Prefer a hand-authored figure in docs/figures/, else the generated one.

    If `placeholder` is given and neither file exists, emit a visible TO-BE-INSERTED
    box instead of failing the build.
    """
    authored = os.path.join(FIGURES_DIR, png)
    generated = os.path.join(HERE, png)
    src = authored if os.path.exists(authored) else generated
    if not os.path.exists(src):
        if placeholder is None:
            raise FileNotFoundError(src)
        ph = doc.add_paragraph()
        ph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        ph.paragraph_format.space_before = Pt(18)
        ph.paragraph_format.space_after = Pt(4)
        r = ph.add_run(placeholder)
        r.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(0xC0, 0x50, 0x20)
        ph2 = doc.add_paragraph()
        ph2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        ph2.paragraph_format.space_after = Pt(18)
        r2 = ph2.add_run("Drop the artwork at docs/figures/" + png
                         + " and rebuild. Build spec: docs/figures/FIGURE-1-SPEC.md")
        r2.italic = True
        r2.font.size = Pt(9)
        r2.font.color.rgb = RGBColor(0x8a, 0x81, 0x78)
    else:
        doc.add_picture(src, width=Inches(width))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(caption)
    r.italic = True
    r.font.size = Pt(9)


# ==========================================================================
# Cover
# ==========================================================================
title("Detailed Design - Bagheera")
para(("GUI System for Managing and Visualizing the PANTHER Model", "b"), space_after=2)
para(("Revision 2 - post-implementation ", "b"),
     ("(supersedes the initial detailed design; reflects the system as actually built)", "i"),
     space_after=10)

para(("Course:", "b"), " Final Project", space_after=2)
para(("Organizational Partner:", "b"), " Sheba Medical Center", space_after=2)
para(("Supervisors:", "b"), " Dr. Sharon Yalov Handzel", space_after=2)
para(("Project Team:", "b"), space_after=2)
para("Dolfin Varshev 315853101, Daniel Rubinstein 314621467, Shahar Ishay 322854308",
     space_after=2)
para(("Repository:", "b"), " github.com/shaharishay14/Bagheera  -  39 commits, 10 April 2026 to 23 July 2026, "
     "~11 200 lines of Python and ~8 900 lines of TypeScript/TSX.", space_after=12)

# ==========================================================================
h("0. Purpose of this revision, and what changed since the initial design", before=8)
# ==========================================================================
para("The first detailed design was written before implementation began. It described a system that "
     "would wrap the PANTHER model in a GUI, and it deliberately left several decisions open, marked "
     "[PENDING SHEBA COORDINATION], because they depended on infrastructure the team had not yet seen.")
para("The system has since been built, containerized, and deployed on a GPU workstation. This revision "
     "replaces every planned or provisional element with what was actually implemented. The substantive "
     "changes are listed below; the rest of the document is the design as built.")

table([
    ["Area", "Initial design (rev. 1)", "As built (rev. 2)"],
    ["Scope",
     "Wrap PANTHER inference only. Feature extraction was assumed to already exist.",
     "The scope grew to the FULL pipeline. Bagheera also drives TRIDENT (tissue segmentation, "
      "tiling, patch encoding), because without it a user still needed a terminal to produce "
      "the feature files PANTHER consumes."],
    ["Central operation",
     "'Execute model inference' - run the trained PANTHER model on a dataset.",
     "'Train a model' is the central operation. The user trains their own PANTHER prototype model "
      "from the browser; running it on new slides became a secondary flow (still present, as UC9)."],
    ["Model organisation",
     "Not specified.",
     "Two generations. K-fold cross-validation was implemented first (one submission -> K models "
      "grouped by a Model Group), then replaced: PANTHER's prototype fit is unsupervised and has "
      "no held-out label to validate against, so K-fold added cost without a payoff. New runs "
      "train ONE standalone model on all slides; existing K-fold groups remain readable (legacy, "
      "read-only). Cross-model inspection moved to a dedicated Model Comparison page."],
    ["Worker architecture",
     "TWO worker processes: an 'Execution Worker' that fires jobs and immediately frees itself, "
      "and a parallel 'Result Worker' that polls running processes for completion.",
     "ONE worker thread that runs each job to completion synchronously. The two-worker split was "
      "dropped: with a single GPU there is nothing to gain from firing jobs in parallel, and one "
      "thread that owns a job start-to-finish removes an entire class of orphaned-process bugs. "
      "The thread never crashes - every handler exception is caught, logged and recorded."],
    ["Queue capability",
     "FIFO with a priority column; drag-to-reorder listed as a use case.",
     "Delivered in full, plus more: cancel a waiting job, drag-to-reorder, and re-run a "
      "failed/cancelled/succeeded job at the tail. Correctness rests on compare-and-swap updates "
      "rather than application locks (section 5.1)."],
    ["Database",
     "Integrate with Sheba's existing relational schema; exact mapping pending.",
     "A self-contained SQLite database (11 tables) owned by the application. No hospital-schema "
      "integration was required for the deliverable. SQLite runs in WAL mode with a 5 s busy "
      "timeout so the HTTP threads and the worker thread share it safely."],
    ["Heavy-file storage",
     "Pending: local network drives, hospital object storage, or PACS.",
     "Plain bind-mounted directories on the GPU host, split across two disks. WSIs are mounted "
      "read-only; all writable state (database, renders, job logs, split CSVs, inference outputs, "
      "model weights cache) lives on mounted volumes outside the container image."],
    ["REST API",
     "Six endpoints sketched with abstracted identifier types.",
     "About 58 endpoints across 14 routers, with concrete Pydantic request/response models. "
      "Identifiers are UUID strings throughout (section 2.3)."],
    ["Visualization",
     "'Cluster visualization' - display prototypes and representative patches.",
     "Four analysis panels per model, reproducing the figure layout of the PANTHER paper: "
      "Section A (whole-slide prototype assignment map, mixture weights, region of interest), "
      "Section B (validation consistency), Section C (on-tissue 2-D embedding map), and "
      "Section D (prototype dictionary). Plus per-slide violin plots and zoom/pan viewers."],
    ["Deployment",
     "On-premise, served over the hospital intranet to concurrent pathologists.",
     "On-premise via Docker Compose on a single GPU workstation, reached over the LAN. "
      "Multi-user concurrency was NOT built - the design remains single-user, single-GPU "
      "(section 12)."],
    ["Testing",
     "Planned. A trace-table template was provided with four example rows.",
     "78 automated tests run green: 50 pytest tests on the backend, 28 Vitest tests on the "
      "frontend (section 8)."],
], widths=[1.0, 2.6, 3.3], size=8.5)

para(("A note on honesty. ", "b"),
     "The renderers in ",
     ("services/visualization.py", "c"),
     " and the inference plumbing in ",
     ("services/inference.py", "c"),
     " were written against upstream source and a reference notebook. Each module opens with a "
     "numbered list of assumptions about PANTHER's prototype file structure, TRIDENT's output "
     "paths and HDF5 key conventions. Several of these assumptions were wrong on first contact "
     "with the real GPU box and were fixed in place (the split directory PANTHER actually reads "
     "from, the directory name it asserts on, the feature width per encoder, the true patch pitch "
     "used when tiling). The remaining unverified assumption is the exact subdirectory TRIDENT "
     "writes slide thumbnails to; the code probes three extensions, falls back to generating the "
     "thumbnail itself, and reports a ",
     ("thumbnails_found", "c"),
     " diagnostic so a mismatch is visible rather than silent.")

doc.add_page_break()

# ==========================================================================
h("1. Project Scope and Requirements", before=0)
# ==========================================================================
para("Bagheera is an interface layer that bridges the gap between advanced computational-pathology "
     "models and the people who need to use them. It provides a graphical user interface for the "
     "PANTHER model and for the TRIDENT feature-extraction pipeline it depends on.")
para("PANTHER performs unsupervised slide representation learning: it discovers a set of ",
     ("prototypes", "i"),
     " - recurring morphological patterns - across a cohort of whole-slide images, and describes "
     "each slide as a mixture over those prototypes. Both PANTHER and TRIDENT ship as command-line "
     "research code. Bagheera turns that raw capability into a navigable visual workspace, so a "
     "clinician or researcher can run the pipeline, see what the model learned, name it in "
     "biological terms, and apply it to new slides without touching a terminal.")

para(("Functional requirements.", "b"),
     " Each requirement states what the system shall do, together with the acceptance criteria "
     "used to decide whether it does. All nine are implemented and exercised by the use cases in "
     "section 6.", space_before=6)

table([
    ["#", "Requirement", "Acceptance criteria"],
    ["FR1", "The system shall extract patch-level features from a directory of whole-slide images "
            "without the user issuing any command-line invocation.",
     "The user selects a slide directory, a dataset name and a patch encoder. The system shall "
     "segment tissue, tile it at 20x magnification into patches of the size the chosen encoder "
     "requires, and write one feature file per slide. The exact command issued shall be shown "
     "before submission and stored with the run."],
    ["FR2", "The system shall generate a reproducible partition of a cohort from a manifest CSV.",
     "Given a CSV and a seed, the system shall write the partition to disk and record the seed, "
     "the row count and the detected slide-identifier column. The same CSV and seed shall always "
     "produce the same partition. A CSV with no slide-identifier column shall be rejected before "
     "anything is written."],
    ["FR3", "The system shall train a PANTHER model from the browser without blocking the user "
            "interface.",
     "The user sets the hyperparameters and submits. The request shall return before any GPU work "
     "begins, the run shall proceed asynchronously, and the user shall be able to observe its "
     "progress and read its log while it runs."],
    ["FR4", "The system shall cluster the tissue of a cohort into a user-specified number of "
            "morphological prototypes, and shall assign every extracted patch of every slide to "
            "exactly one prototype, without requiring any manual annotation or label.",
     "(a) The user shall specify the number of prototypes and the clustering backend (K-means, or "
     "FAISS on the GPU) before submission; both shall be recorded on the trained model. "
     "(b) Training input shall consist only of the feature files and the split manifest - no "
     "labels, no regions drawn by hand. "
     "(c) On completion the system shall persist the fitted prototype centres, and shall be able "
     "to produce, for any slide of the cohort, a prototype assignment for each of its patches and "
     "a set of per-slide mixture weights over the prototypes that sum to one. "
     "(d) Each prototype shall carry a stable index, presented as C1..Cn, which every downstream "
     "view, colour and user-assigned label refers to. "
     "(e) Given the same features, the same split and the same seed, the fit shall be reproducible."],
    ["FR5", "The system shall serialise all heavy work through a single shared queue that the user "
            "can inspect and control.",
     "The queue shall show what is running, what is waiting and in what order, and recently "
     "finished work. The user shall be able to cancel a waiting job, change the waiting order by "
     "dragging, and re-run a finished or failed job. A running job shall not be cancellable "
     "(section 12)."],
    ["FR6", "The system shall present the clustering of FR4 in a form a pathologist can interpret "
            "without reading code or opening a file.",
     "For a trained model the system shall render: a map of the slide with every patch painted in "
     "its prototype's colour; a chart of the proportion of tissue each prototype accounts for; a "
     "dictionary showing the most representative patches of each prototype; and a two-dimensional "
     "embedding showing how prototypes relate to one another. Every image shall be zoomable and "
     "pannable, and the magnified region of interest shall be re-selectable, either automatically "
     "or by clicking a point on the slide."],
    ["FR7", "The system shall allow the clustering produced by different models to be compared on "
            "the same tissue.",
     "The user shall be able to render one slide through up to four models side by side, and "
     "shall be able to lock zoom and pan across the columns so the same tissue is viewed at the "
     "same magnification in each."],
    ["FR8", "The system shall let a domain expert attach meaning to what the model found.",
     "The user shall be able to assign a biological name to each prototype, and to attach and "
     "edit free-text notes on a model and on an individual inference result. A prototype shall "
     "carry at most one name, and names shall persist across sessions."],
    ["FR9", "The system shall apply a trained model to slides that were not in its training "
            "cohort.",
     "The user shall be able to submit one or many new slides. The system shall extract their "
     "features using the same encoder configuration the model was trained against, assign their "
     "patches to that model's prototypes, and render the result. A slide already processed by "
     "that model shall be served from cache rather than recomputed, with an explicit per-slide "
     "override to force a re-run."],
], widths=[0.45, 2.15, 4.3], size=8)

para(("Non-functional requirements.", "b"),
     " Each is stated with the criterion used to check it and the value measured on the delivered "
     "system, so the requirement can be re-verified rather than argued about.", space_before=8)

table([
    ["#", "Requirement", "Measurable criterion", "Measured / verified"],
    ["NFR1", "Data locality",
     "The backend initiates zero outbound network connections during normal operation. No patient "
     "data leaves the host under any code path.",
     "Met. All model weights are fetched at build time or on first use and cached in a mounted "
     "volume; no run-time code path opens an external connection. One caveat, stated for honesty: "
     "the browser page requests two web fonts from Google Fonts. The request carries no "
     "application data, and the interface falls back to system fonts when it fails, but on an "
     "air-gapped network the fonts should be vendored into the image."],
    ["NFR2", "Development without a GPU",
     "The API starts and answers GET /api/health with 200 on a machine where torch, openslide, "
     "h5py, matplotlib and umap are NOT installed, in under 2 seconds.",
     "Met, measured on a laptop: 550 ms to import the application, 27 ms for startup (schema "
     "creation, storage directories, handler registration, worker thread) plus the first request. "
     "Total ~0.6 s. Confirmed at run time that torch, openslide and matplotlib were never loaded "
     "into the process."],
    ["NFR3", "Fault isolation",
     "No failure inside a job - non-zero subprocess exit, missing dependency, malformed input, "
     "unhandled exception - may terminate the worker thread or make the API unresponsive. The "
     "next queued job must start within one poll interval (2 s).",
     "Met. Every handler runs inside a broad catch that records the traceback in the job log and "
     "the message on the job row, then continues the loop. Verified in production use: four "
     "missing machine-learning dependencies and three upstream path assertions were each "
     "diagnosed and fixed while the application stayed up."],
    ["NFR4", "Path confinement",
     "100% of endpoints that accept a user-supplied path resolve it through the single sandbox "
     "function before any filesystem operation. Traversal and symlink escape both yield 403.",
     "Met: 28 call sites across 8 routers and 3 services; there is no second code path that "
     "opens a user-supplied file. Symlinks are collapsed by resolution before the check, so a "
     "symlink pointing outside the allow-list is rejected. Regression test: "
     "test_403_outside_roots."],
    ["NFR5", "Submission responsiveness",
     "Submitting long-running work must return in well under 1 second regardless of dataset "
     "size, and the interface must reflect a state change within 2 seconds of it happening.",
     "Met by construction. A training submission performs two inserts and returns; no GPU work "
     "happens in the request. All progress surfaces poll on a 2 s cadence, so worst-case "
     "staleness is one interval. The one deliberate exception is the initial TRIDENT extraction, "
     "which is synchronous by design (section 12)."],
    ["NFR6", "Durability across restart",
     "Stopping and rebuilding the containers must lose no completed work and no user-entered "
     "data. A full backup must be a file copy.",
     "Met. The image holds no state: the database, rendered figures, job logs, split CSVs, "
     "inference outputs and cached model weights all live on bind-mounted host directories. "
     "Backup is copying one SQLite file plus the output directories. The one gap is that a job "
     "which was RUNNING at shutdown does not resume - it must be re-run from the queue page."],
    ["NFR7", "Reproducibility",
     "The same inputs and the same seed must produce the same split, the same preview slide "
     "selection and the same region of interest, on any machine.",
     "Met. Splits are generated with a seeded random source and record their seed; preview-slide "
     "selection is seeded by the model identity; region-of-interest windows are ranked "
     "deterministically. The exact command issued to each pipeline is stored with the run, so any "
     "result can be traced back to the invocation that produced it."],
    ["NFR8", "Portability of deployment",
     "A new GPU machine must reach a working installation with no manual dependency resolution: "
     "clone, edit one environment file, run two commands. The same repository must build against "
     "more than one CUDA generation.",
     "Met. docker compose build and docker compose up, with the CUDA base image, the PyTorch "
     "wheel index and the FAISS build all exposed as build arguments - CUDA 12.1 and 11.8 are "
     "both supported by changing two values."],
    ["NFR9", "Maintainability and testability",
     "The full automated suite must run on a developer laptop, with no GPU and no slides, in "
     "under 10 seconds - fast enough to run on every change.",
     "Met: 78 tests (50 backend, 28 frontend) in roughly 2.6 seconds combined. Each backend test "
     "gets its own temporary database; no test loads a machine-learning library, touches a real "
     "slide or spawns a subprocess."],
    ["NFR10", "Operability without a terminal",
     "Every one of the thirteen use cases must be completable from the browser, and every "
     "long-running operation must expose its live log in the interface.",
     "Met for the twelve user-facing use cases; UC13 (deployment) is an operator task and is "
     "expected to use a shell. Every asynchronous job writes a log file whose tail is reachable "
     "from every surface that shows that job's status."],
], widths=[0.45, 1.25, 2.4, 2.8], size=8)

doc.add_page_break()

# ==========================================================================
h("2. Logical Architecture and Modular Decomposition", before=0)
# ==========================================================================
para("The system is a client-server application decomposed into layers with a single direction of "
     "dependency. Nothing in a lower layer knows about a higher one.")

figure("fig1_architecture.png",
       "Figure 1 - Logical architecture and physical deployment. Arrows show dependency and data "
       "flow; each shaded region is a separate deployment unit.",
       placeholder="[ FIGURE 1 - TO BE INSERTED ]")

h("2.1 Core modules", before=8, after=4)
bullet("Client (view). ",
       "A Vite + React 18 + TypeScript single-page application styled with Tailwind. It holds no "
       "business logic: it renders server state, submits forms, and polls. Ten page components "
       "behind eleven routes.")
bullet("API layer (controller). ",
       "FastAPI. Fourteen routers, one per resource, each thin: validate the request, call a "
       "service, shape the response. Routers never touch the filesystem or a subprocess directly.")
bullet("Service layer (model). ",
       "Plain Python modules holding all business logic - the path sandbox, the split generator, "
       "the TRIDENT and PANTHER subprocess builders, the renderers, the caches, the deletion "
       "cascade. Services do not know that HTTP exists.")
bullet("Worker. ",
       "One background daemon thread that drains the job table and invokes registered handlers. "
       "The handlers are what actually shell out to the two research repositories on the GPU.")
bullet("Persistence. ",
       "SQLite for structured metadata; the filesystem for everything heavy (whole-slide images, "
       "feature files, prototype files, rendered images, job logs).")

h("2.2 Cohesion, coupling, and the interface contract", before=10, after=4)
para("The rev. 1 design argued that the API and the model worker should be decoupled through the "
     "database rather than by calling each other. That decision was kept, and it proved to be the "
     "single most valuable structural choice in the project.")
bullet("High cohesion. ",
       "The API is responsible for routing, validation and persistence, and performs no heavy "
       "computation. The worker is dedicated to executing pipelines and rendering figures, and is "
       "entirely unaware of HTTP requests, sessions or responses.")
bullet("Loose coupling through the jobs table. ",
       "The API and the worker never call one another. The API inserts a row describing the work "
       "to be done; the worker reads it. In practice this meant that when a training run died "
       "because a CUDA build of FAISS was missing, or because PANTHER asserted on a directory "
       "name, the browser stayed fully responsive and the failure surfaced as a red status pill "
       "with a readable log - not as a hung request.")
bullet("A second contract on the client side. ",
       "The file lib/api.ts is the single source of truth for the frontend-backend contract: one "
       "TypeScript interface per response shape and one function per endpoint. Pages import from "
       "it and never call fetch directly, so changing an endpoint without changing this file is a "
       "compile error rather than a runtime surprise.")

para(("The polymorphic job pointer. ", "b"),
     "A job row does not know what kind of object it operates on. It carries a pair - ",
     ("ref_table", "c"), " and ", ("ref_id", "c"),
     " - plus an optional JSON parameter blob. This is what let four different job types "
     "(training, post-training visualization, on-demand slide rendering, inference) share one "
     "queue, one worker and one status UI without any of them knowing about the others.",
     space_before=6)

h("2.3 REST API specification", before=10, after=4)
para("All endpoints are served under ", ("/api", "c"),
     ". Identifiers are UUID strings. Request and response bodies are Pydantic models, so the "
     "interactive OpenAPI specification at ", ("/docs", "c"),
     " is generated from the same definitions the server validates against. The table below groups "
     "the surface by resource; representative endpoints are shown rather than all 58.")

table([
    ["Router", "Representative endpoints", "Purpose"],
    ["fs",
     "GET /fs/roots · GET /fs/list · GET /fs/csv-count · GET /fs/csv-inspect",
     "Sandboxed directory browsing and CSV diagnostics. Every path passes the allow-list check."],
    ["trident, runs",
     "POST /trident/run · GET /trident/runs · GET /runs/resolve",
     "Launch feature extraction (synchronous), list past runs, and map a features directory back "
     "to the run that produced it."],
    ["splits",
     "POST /splits · GET /splits · GET /splits/{id}",
     "Create a split from a CSV. The body carries kind = 'single' (100% train, the default for new "
     "runs) or 'kfold' (legacy)."],
    ["panther",
     "POST /panther/single-runs · POST /panther/runs · GET /panther/models",
     "Train one standalone model (returns model_id + job_id immediately), or the legacy K-fold "
     "submission. Listing supports a run_kind filter."],
    ["models",
     "GET /models/{id} · PATCH /models/{id} · POST /models/{id}/repick-roi · "
     "POST /models/{id}/select-roi · POST /models/{id}/render-slide · GET /models/{id}/slide-viz · "
     "DELETE /model-groups/{id}",
     "Inspect and edit a model, drive the region-of-interest controls, request an on-demand "
     "per-slide render for the comparison page, and cascade-delete a legacy K-fold group with "
     "its artifacts."],
    ["datasets, thumbnails",
     "GET /datasets · GET /datasets/{name}/slides · GET /datasets/{name}/models · "
     "GET /slide-thumbnail",
     "Backing lookups for the Model Comparison page: which datasets have models, which slides "
     "exist, and a binary thumbnail per slide for the picker grid."],
    ["labels, notes",
     "GET/POST/DELETE /prototype-labels · GET/POST/PATCH/DELETE /model-notes",
     "Semantic labelling of prototypes and free-text notes on a model."],
    ["inference",
     "POST /inferences · GET /inferences/lookup · GET /inferences/{id}/example-patches · "
     "POST /inferences/{id}/rerun · GET/POST/PATCH/DELETE /inference-notes",
     "Single and batch inference dispatch with a cache pre-check, per-result example patches, "
     "and notes."],
    ["jobs, queue",
     "GET /jobs · GET /jobs/{id} · POST /jobs/{id}/cancel · POST /jobs/{id}/retry · "
     "GET /queue · POST /queue/reorder",
     "The queue surface: list and inspect jobs (including a tail of the live log), cancel a "
     "waiting job, re-run a terminal one, and atomically rewrite the waiting order."],
    ["viz",
     "GET /viz/{file_path} · GET /viz/placeholder/{kind}",
     "Serve a rendered image from one of the two output roots (403 outside them, 404 if missing), "
     "or synthesize a labelled SVG placeholder for a render that does not exist yet."],
], widths=[0.85, 2.5, 3.55], size=8.5)

para(("Status codes carry meaning. ", "b"),
     "403 means a path fell outside the allowed roots. 409 signals a lost race or an invalid state "
     "transition - cancelling a job the worker has already claimed, deleting a model group with an "
     "active job, re-picking a region of interest before the panel has been rendered. 422 is a "
     "well-formed request the system cannot satisfy (a slide with no tissue window, a "
     "thumbnail the imaging library could not open). 207 is returned by batch dispatch when some "
     "paths validated and others did not, and the client renders the per-path outcome rather than "
     "treating the whole batch as failed.")

doc.add_page_break()

# ==========================================================================
h("3. Physical Deployment and Infrastructure", before=0)
# ==========================================================================
para("The system is deployed on-premise, inside the hospital's network, with no external internet "
     "access required at run time. It is delivered as two Docker containers orchestrated by a "
     "single Compose file, so the entire installation on a new GPU machine is: clone the "
     "repository, fill in an environment file, and run two commands.")

table([
    ["Container", "Contents", "Exposed"],
    ["frontend",
     "nginx serving the built single-page application, and reverse-proxying /api/* to the backend. "
     "Nothing else - no Node process at run time.",
     "Port 8080 on the host (configurable)."],
    ["backend",
     "A CUDA runtime base image with one shared Python virtual environment containing FastAPI, the "
     "worker, PyTorch, OpenSlide, a GPU build of FAISS, and pinned checkouts of the TRIDENT and "
     "PANTHER repositories baked into the image.",
     "Not published; reached only through the frontend proxy."],
], widths=[0.9, 4.4, 1.6], size=8.5)

para(("GPU access", "b"),
     " is granted through the NVIDIA Container Toolkit and a device reservation in the Compose "
     "file. The CUDA base image and the PyTorch wheel index are build arguments, so the same "
     "repository builds against either a CUDA 12.1 or a CUDA 11.8 driver by changing two lines in "
     "the environment file.")

para(("Storage layout. ", "b"),
     "The container image is disposable; all data and state live on bind-mounted host directories, "
     "split across two physical disks because the render and inference outputs grow much faster "
     "than the database:", space_before=6)

table([
    ["Host location", "Mounted at", "Holds"],
    ["WSI data directory (read-only)", "/data", "Whole-slide images and the manifest CSV. Mounted "
     "read-only so no code path can modify the source cohort."],
    ["Internal disk / db", "/state/db", "bagheera.db - the SQLite database."],
    ["Internal disk / hf_cache", "/state/hf_cache", "Downloaded encoder weights, so a restart does "
     "not re-download them."],
    ["External disk / viz_cache", "/state/viz_cache", "All rendered figures, plus one log file per "
     "job under job_logs/."],
    ["External disk / inference_outputs", "/state/inference_outputs",
     "Per-inference feature files and rendered figures."],
    ["External disk / datasets_splits", "PANTHER's splits directory",
     "The generated train/val/test CSVs, mounted into the exact location PANTHER reads them from."],
    ["External disk / trident_processed", "/app/backend/trident_processed",
     "TRIDENT feature-extraction outputs for the initial dataset build."],
], widths=[1.85, 1.85, 3.2], size=8.5)

para(("A deployment lesson worth recording. ", "b"),
     "The path model - which paths are host paths, which are container paths, and which of them "
     "the file browser is allowed to traverse - was the single largest source of first-deploy "
     "friction. The allow-list must name ",
     ("container", "i"),
     " paths, and it must include the TRIDENT output directory as well as the data directory, "
     "because the training form needs to browse to the features it just produced. A separate "
     "document (",
     ("docs/env-paths.md", "c"),
     ") was written for this reason alone. A related failure: the external disk is a network share "
     "that returns a permission error on symlink creation, so a compatibility symlink PANTHER "
     "requires is now created on the container's local filesystem with an absolute target instead.")

doc.add_page_break()

# ==========================================================================
h("4. Data Architecture and Schema Design", before=0)
# ==========================================================================
para("The system enforces a strict separation between structured metadata and heavy computational "
     "data. Metadata - runs, splits, models, labels, notes, jobs - lives in a relational database. "
     "Whole-slide images, HDF5 feature files, prototype files and rendered figures live on the "
     "filesystem, and the database stores only their paths. This keeps the database small enough "
     "that it is trivially backed up by copying one file.")

h("4.1 Relational database", before=8, after=4)
para("SQLite, accessed through SQLAlchemy 2.0 with typed models. Eleven tables:")

table([
    ["Table", "One row per...", "Notable columns and constraints"],
    ["trident_runs", "one feature-extraction run",
     "dataset_name, wsi_dir, patch_encoder, mag, patch_size, the exact command executed, status, "
     "captured stdout/stderr, output_dir."],
    ["splits", "one split written to disk",
     "split_name (unique), abs_path, source_csv, k, seed, total_rows, per_fold_counts (JSON)."],
    ["model_groups", "one legacy K-fold submission",
     "Legacy only - no new rows are created. Retained so existing groups stay readable."],
    ["models", "one trained model",
     "The central row. Naming, grouping (group_id, fold_index, fold_k, run_kind), inputs "
     "(features_dir, split, trident_run), PANTHER hyperparameters (mode, in_dim, n_proto, "
     "n_proto_patches, n_init, seed), outcome (status, prototypes_dir), and rendered artifacts "
     "(viz_status, viz_artifacts JSON). Unique on (group_id, fold_index)."],
    ["panther_runs", "one training subprocess attempt",
     "An execution log: the command, status, return code, stdout, stderr, start and finish times. "
     "Kept even when the model row is updated, so a failure is always inspectable."],
    ["prototype_labels", "one label on one prototype",
     "Unique on (model_id, prototype_index) - a prototype has at most one name."],
    ["model_notes", "one free-text note on a model", "body, created_at, updated_at."],
    ["inference_batches", "a group of inferences submitted together", "model_id, user_label, total_count."],
    ["inferences", "one (model, slide) run",
     "Slide identity (path, filename, mtime, size, sha256 hash), output paths, status. "
     "Unique on (model_id, wsi_hash) - this pair IS the cache key."],
    ["inference_notes", "one free-text note on an inference result", "body, created_at, updated_at."],
    ["jobs", "one asynchronous unit of work",
     "job_type, the polymorphic (ref_table, ref_id) pointer, status, queue_position, log_path, "
     "error_message, and an optional params JSON blob."],
], widths=[1.15, 1.85, 3.9], size=8.5)

para(("Schema lifecycle - a deliberate constraint. ", "b"),
     "There is no migration framework. The schema is created by ",
     ("Base.metadata.create_all()", "c"),
     " at startup, and after a structural change the developer deletes the database file and lets "
     "it recreate. This is acceptable because the database holds derived bookkeeping, not primary "
     "clinical data. For the cases where an existing deployment's data had to survive - adding ",
     ("queue_position", "c"), ", ", ("viz_artifacts", "c"), ", ", ("run_kind", "c"), " and ",
     ("params", "c"),
     " - the convention is a guarded ", ("ALTER TABLE ... ADD COLUMN", "c"),
     " executed inside the startup routine. Four such guards exist today, each documented in the "
     "database reference. This gives the benefit of additive migrations without the weight of a "
     "framework.", space_before=6)

para(("Concurrency. ", "b"),
     "The database is shared between the HTTP request threads and the worker thread. It runs in "
     "write-ahead-logging mode with a five-second busy timeout, which is what keeps many "
     "concurrent pollers and one occasional writer from producing 'database is locked' errors.")

h("4.2 Filesystem layout", before=10, after=4)
code([
    "${VIZ_CACHE_ROOT}/",
    "  job_logs/{job_id}.log                one append-only log per asynchronous job",
    "  slide_thumbs/{key}.jpg               generated slide-picker thumbnails (content-keyed)",
    "  {model_id}/",
    "    section_a/ section_b/ section_c/ section_d/    the four analysis panels",
    "    section_b_violin/                  per-slide violin plot",
    "    umap.png                           dataset-wide embedding scatter",
    "    compare/{slide_id}/                on-demand per-slide render + manifest.json",
    "",
    "${INFERENCE_ROOT}/",
    "  {model_id}/{inference_id}/",
    "    custom_list.csv                    the single-slide manifest handed to TRIDENT",
    "    trident_output/.../{slide}.h5      the extracted features",
    "    heatmap_{slide}.png  mixture_{slide}.png  tsne_{slide}.png",
    "    example_patches_{slide}/prototype_{NN}/patch_{NN}.png",
])
para("Rendered artifacts are addressed by absolute path in the database and served through a single "
     "endpoint that re-checks the path against both output roots on every request. Both roots grow "
     "without bound; pruning is a manual operation today and is listed as known work in section 12.")

doc.add_page_break()

# ==========================================================================
h("5. Algorithm Descriptions", before=0)
# ==========================================================================
para("Four algorithms carry most of the system's non-obvious behaviour: the queue, the split "
     "generator, the region-of-interest selection, and the two caches.")

h("5.1 Queue management", before=8, after=4)
para(("Plain-English description.", "b"))
numbered("1. The jobs table is the queue. A new job is inserted with status 'queued' and a "
         "queue_position one greater than the current maximum, which yields first-in-first-out "
         "ordering by default.")
numbered("2. A single worker thread wakes every two seconds and selects the queued job with the "
         "lowest queue_position.")
numbered("3. It claims that job with a conditional update: set status to 'running' only where the "
         "status is still 'queued'. If zero rows change, a user cancelled the job in the "
         "meantime, so the worker moves on to the next one.")
numbered("4. On claiming, queue_position is cleared - the job has left the waiting line.")
numbered("5. The worker dispatches to the handler registered for the job's type and runs it to "
         "completion. Everything the handler prints is streamed to a per-job log file.")
numbered("6. On return, the job is marked 'succeeded'. On any exception, the traceback is written "
         "to the log and the job is marked 'failed' with the error message stored on the row. "
         "The loop continues either way; the worker never dies.")
numbered("7. A handler may enqueue further work. Training enqueues a visualization job on success, "
         "which re-enters the same queue at the tail.")
numbered("8. A user may cancel a waiting job (the same conditional update in reverse), drag to "
         "reorder the waiting jobs (an atomic rewrite of queue_position for the whole waiting "
         "set), or re-run a terminal job (which re-enters at the tail).")

para(("Pseudocode.", "b"), space_before=6)
code([
    "// The worker loop - one thread, one job at a time.",
    "function workerLoop(stopEvent):",
    "    while not stopEvent.isSet():",
    "        job = claimNextJob()",
    "        if job is null:",
    "            wait(POLL_INTERVAL = 2s); continue",
    "        runOneJob(job)",
    "",
    "function claimNextJob():",
    "    loop:",
    "        job = SELECT * FROM jobs WHERE status='queued'",
    "                ORDER BY queue_position, created_at LIMIT 1",
    "        if job is null: return null",
    "        // compare-and-swap: the database decides who wins the race",
    "        n = UPDATE jobs SET status='running', started_at=now(), queue_position=NULL",
    "              WHERE id=job.id AND status='queued'",
    "        if n == 1: return job",
    "        // n == 0 -> cancelled a moment ago; try the next one",
    "",
    "function runOneJob(job):",
    "    log = openLog(job.logPath)",
    "    try:",
    "        handler = HANDLERS[job.jobType]      // panther_train | post_train_viz",
    "        handler(db, job, log)                //   | render_slide | inference",
    "        job.status = 'succeeded'",
    "    catch anyError as e:                     // the worker must never crash",
    "        log.write(traceback())",
    "        job.status = 'failed'; job.errorMessage = truncate(e, 2000)",
    "    finally:",
    "        job.finishedAt = now(); log.close()",
    "",
    "// User-initiated transitions, both compare-and-swap.",
    "function cancel(jobId):",
    "    n = UPDATE jobs SET status='canceled', finished_at=now(), queue_position=NULL",
    "          WHERE id=jobId AND status='queued'",
    "    if n == 0: return HTTP 409   // the worker already started it",
    "",
    "function reorder(orderedIds):                // atomic, in one transaction",
    "    queued = SELECT * FROM jobs WHERE status='queued'",
    "    pos = 1",
    "    for id in orderedIds:  if id in queued: queued[id].position = pos++",
    "    for job in queued not yet assigned:      // keep their relative order at the back",
    "        job.position = pos++",
])

figure("fig2_sequence.png",
       "Figure 2 - Sequence diagram for a PANTHER training submission. The HTTP request completes "
       "before any GPU work begins; the jobs table is the only channel between the API and the worker.")

para(("Why this differs from the rev. 1 algorithm. ", "b"),
     "The original design proposed two cooperating workers: an execution worker that fired a job "
     "asynchronously and immediately freed itself, and a result worker that polled the operating "
     "system for process completion. That design serves a machine that can run several jobs at "
     "once. With one GPU it buys nothing and costs a great deal: process handles must be tracked "
     "across restarts, a crashed subprocess leaves a job permanently 'processing', and the two "
     "workers race over the same rows. Running each job to completion inside the claiming thread "
     "removes all of it. The cost is stated plainly in section 12: a running job cannot be "
     "interrupted.")

para(("Complexity.", "b"), space_before=6)
bullet("Claiming a job: ", "one indexed SELECT plus one conditional UPDATE, so O(log n) in the "
       "number of queued rows. The retry loop iterates only when a cancellation wins a race, which "
       "is bounded by the number of cancellations, not by queue length.")
bullet("Reordering: ", "O(m) for m waiting jobs, in a single transaction.")
bullet("Building the queue view: ", "O(w + r) for w waiting jobs and r recent jobs, where r is "
       "capped at 20.")
bullet("Memory: ", "O(1) auxiliary. The queue holds job metadata only - identifiers, statuses, "
       "paths. Feature vectors and images are never loaded into the queue layer; the handler reads "
       "them, and they are released when the job ends.")

h("5.2 Split generation", before=10, after=4)
para("A split turns a manifest CSV into the directory structure PANTHER expects. Two shapes exist.")
bullet("Single split (the default for new runs). ",
       "One fold containing every slide in train.csv, with header-only val.csv and test.csv so "
       "that any downstream glob still resolves. The split is named "
       "alltrain_seed_{seed}_{random8}, and a metadata file records the seed, the row count and "
       "the detected slide-identifier column.")
bullet("K-fold split (legacy). ",
       "Rows are shuffled with a seeded generator and chunked into K parts; fold i uses chunk i as "
       "test, chunk (i+1) mod K as validation, and the rest as train. The invariant is that every "
       "slide appears in exactly one fold's test set.")
para(("Input validation that was added after first contact with real data. ", "b"),
     "The slide-identifier column is detected case-insensitively against a known set of names. If "
     "no such column exists, the request is rejected with a 400 rather than writing a manifest the "
     "model cannot read. If it exists, a trailing .tif or .tiff is stripped from each value before "
     "the generated CSVs are written - the source CSV on the read-only mount is never modified. "
     "This single normalisation removed an entire class of silent training failures.")

h("5.3 Region-of-interest selection", before=10, after=4)
para("The lead analysis panel reproduces the PANTHER paper's figure: a magnified square region of "
     "the slide shown twice, once as raw tissue and once tiled with each patch tinted by the "
     "prototype it was assigned to. Choosing which region to show is not obvious, and the "
     "algorithm went through several corrections against real slides.")
bullet("Candidate windows. ",
       "Non-overlapping windows of 16 x 16 patches are enumerated over the slide's patch "
       "coordinates and ranked by occupancy (how many grid cells actually contain an extracted "
       "patch), then by prototype diversity, then by density. Ranking by diversity first was the "
       "original approach and it consistently landed on tissue boundaries that were half "
       "background; occupancy-first lands in the middle of dense tissue and produces the packed "
       "tiling the paper shows.")
bullet("The true patch pitch. ",
       "Patches are not necessarily laid out on the stored patch size. A dataset that extracts "
       "256-pixel patches at half the base magnification places them on a 512-pixel grid at "
       "level 0. Tiling on the stored size therefore painted only every other cell. The pitch is "
       "now derived as the statistical mode of the gaps between adjacent coordinates, which is "
       "robust to stray off-grid patches.")
bullet("Framing. ",
       "The tissue bounding box is computed from the 1st and 99th percentiles of the patch "
       "coordinates rather than their raw minimum and maximum, so a handful of artifact patches "
       "in the corners of the glass cannot stretch the default view across empty slide.")
bullet("Two user controls. ",
       "'Re-pick' advances to the next-ranked window, wrapping around. 'Select' arms the image "
       "viewer as a point picker: the user clicks anywhere on the assignment map, the click is "
       "reported as a normalised coordinate pair, and the server renders the window containing "
       "that point - or, if the click landed on a gap, the window whose centre is nearest. Both "
       "operations are synchronous and re-use a cached array of coordinates and assignments, so "
       "neither re-runs the encoder.")

h("5.4 Caching", before=10, after=4)
bullet("Inference cache (persistent). ",
       "A slide is identified by the SHA-256 of its bytes, streamed in one-megabyte chunks. "
       "Modification time and size are checked first as a cheap pre-filter, but the hash is the "
       "source of truth, and (model_id, hash) is a unique constraint in the database. Submitting a "
       "slide that has already been processed by that model returns the existing result instead of "
       "recomputing it; the user can override this per file with a 're-run' toggle.")
bullet("Assignment cache (in-memory). ",
       "Rendering several panels for the same (model, slide) pair would otherwise run the encoder "
       "once per panel. A process-local, thread-safe cache of the last 32 encoder passes, keyed by "
       "model and slide, is consulted first. The key is model-scoped deliberately - each model has "
       "its own prototypes, so assignments must never be shared between models. Locking is "
       "per-key, so a miss on one slide never blocks another, and two concurrent renders of the "
       "same slide compute once. The dataset-wide renderers deliberately bypass this cache: they "
       "stream every slide's features once and would evict everything useful.")

doc.add_page_break()

# ==========================================================================
h("6. Use Case Analysis and System Interactions", before=0)
# ==========================================================================
para("Thirteen use cases were implemented. The rev. 1 document listed six; the growth comes almost "
     "entirely from taking on feature extraction and model training, which rev. 1 assumed would "
     "happen outside the system.")

figure("fig3_usecase.png", "Figure 3 - Use-case model as built.", width=6.0)

table([
    ["ID", "Use case", "Primary flow"],
    ["UC1", "Extract features (TRIDENT)",
     "The user names a dataset, browses to a directory of slides, and picks a patch encoder. "
     "Locked fields (task, magnification, patch size derived from the encoder) are shown read-only "
     "alongside a live preview of the exact command. Submission blocks until extraction finishes."],
    ["UC2", "Create a split",
     "The user browses to a manifest CSV. The system reports its row count and diagnostics "
     "(columns present, whether a slide-identifier column exists, how many values carry a .tif "
     "suffix), then writes a seeded split to disk."],
    ["UC3", "Train a PANTHER model",
     "The user binds a features directory to its extraction run, chooses or creates a split, sets "
     "hyperparameters, and submits. The response is immediate; the browser navigates to the new "
     "model's page and polls while the worker trains it and then renders its panels."],
    ["UC4", "Browse and manage models",
     "A filterable card grid over all trained models, with search, dataset filter, sort, and a "
     "favourites-only toggle. Legacy K-fold groups appear as separate, chip-tagged cards."],
    ["UC5", "Inspect a model",
     "Four analysis panels. Section A: whole-slide thumbnail with a scale bar, the prototype "
     "assignment map, the mixture-weight bar chart, and the region of interest. Section B: "
     "validation consistency (only meaningful for legacy K-fold models, which have held-out "
     "slides). Section C: an on-tissue two-dimensional embedding map beside the abstract scatter. "
     "Section D: the prototype dictionary - one column per prototype showing its most "
     "representative patches in that prototype's colour."],
    ["UC6", "Re-pick or click-select a region of interest",
     "See section 5.3. Both controls update the stored panel and return the new artifact paths."],
    ["UC7", "Label prototypes and take notes",
     "A name box under each prototype column, saved on blur, plus a notes thread on the model and "
     "on each inference result."],
    ["UC8", "Compare models on one slide",
     "A strict top-down gate: choose a dataset, then a slide from a thumbnail grid, then up to "
     "four models. Each column renders that slide through its model. A synchronise toggle lifts a "
     "single zoom/pan transform and feeds it to the matching image in every column, so the same "
     "tissue is compared at the same magnification."],
    ["UC9", "Run inference on new slides",
     "Single or batch. Each selected path is checked against the cache before submission and "
     "marked as cached or to-be-processed, with a per-path re-run override. Results show the "
     "assignment heatmap, mixture plot, example patches with a lightbox, and a t-SNE projection."],
    ["UC10", "Manage the job queue",
     "A shared page showing the running job, the waiting list (draggable, each cancellable) and "
     "recent terminal jobs (each re-runnable). Polling is frozen while a drag is in progress so "
     "the row under the cursor cannot jump."],
    ["UC11", "Read a job log",
     "Every job status surface exposes a log link that opens the tail of that job's log file in a "
     "modal, refreshing while the job is in flight."],
    ["UC12", "Delete a legacy model group",
     "Removes the group's database rows in dependency order within one transaction plus the "
     "on-disk artifacts, refusing with a 409 if a job is currently running against it. Shared "
     "splits and extraction runs are left untouched. Note: this is offered for legacy K-fold "
     "groups only - a standalone single model has no group row and cannot yet be deleted from "
     "the interface (section 12)."],
    ["UC13", "Deploy and configure",
     "The operator fills in the environment file, sets the allow-list of browsable container "
     "paths, and runs the build. Verification steps are documented."],
], widths=[0.45, 1.55, 4.9], size=8.5)

doc.add_page_break()

# ==========================================================================
h("7. Programming Languages, Tools, and Standards", before=0)
# ==========================================================================
table([
    ["Concern", "Choice"],
    ["Backend language", "Python 3.10+ (postponed annotation evaluation used throughout)"],
    ["Web framework", "FastAPI 0.115 on Uvicorn 0.32"],
    ["ORM and database", "SQLAlchemy 2.0 with typed models; SQLite in write-ahead-logging mode"],
    ["Validation", "Pydantic 2.10"],
    ["Frontend", "React 18, TypeScript 5.7 (strict), Vite 5, Tailwind CSS 3.4, React Router 6"],
    ["Frontend data layer", "The platform fetch API only - no Axios, React Query or SWR, and no "
     "state-management library. Local component state and polling."],
    ["Machine learning", "PyTorch, OpenSlide, h5py, scikit-learn, UMAP, Matplotlib, a CUDA build "
     "of FAISS - all imported lazily inside the functions that need them"],
    ["Testing", "pytest with FastAPI's TestClient (backend); Vitest with React Testing Library and "
     "jsdom (frontend)"],
    ["Packaging and deployment", "Docker and Docker Compose; nginx for static serving and proxying"],
    ["Version control", "Git. Feature branches and structured commit messages that state the "
     "problem, the root cause and the fix."],
], widths=[1.5, 5.4], size=9)

para(("Coding standards enforced across the project:", "b"), space_before=8)
bullet("Lazy heavy imports. ",
       "PyTorch, OpenSlide, h5py and the plotting stack are imported inside functions, never at "
       "module scope. The API boots on a machine with no GPU and no machine-learning dependencies "
       "installed, which is what made laptop development and dependency-free tests possible.")
bullet("One security primitive, used everywhere. ",
       "Every route that accepts a path calls the same resolution function. There is no second way "
       "to open a user-supplied file.")
bullet("The worker catches everything. ",
       "No handler exception may escape to the loop.")
bullet("The API client comes first. ",
       "A new endpoint is added to the typed client before any page uses it.")
bullet("Design tokens only. ",
       "Structural UI elements use semantic tokens (surface, ink, border, accent) rather than "
       "hard-coded colours, so the palette was changed once, globally, when the product was "
       "re-branded.")
bullet("Documentation is part of the change. ",
       "Each subsystem owns a reference document, and the convention is that a code change and its "
       "document change land together.")

para(("Tooling note.", "b"), " The repository carries a set of agent definitions under ",
     (".claude/agents/", "c"),
     ", one per subsystem, each bound to the reference document it owns. This was how "
     "documentation was kept synchronised with a fast-moving codebase across a three-month build.",
     space_before=6)

para(("Mentoring and review. ", "b"),
     "The academic supervisor oversaw methodology; the upstream pipelines are the published work "
     "of the Mahmood Lab at Harvard Medical School, and their repositories are the reference "
     "against which the integration was validated.")

doc.add_page_break()

# ==========================================================================
h("8. Testing Framework and Failure Modes", before=0)
# ==========================================================================
para("Testing is split so that system logic can be verified with no GPU, no machine-learning "
     "dependencies and no real slides. All 78 automated tests run in under three seconds on a "
     "laptop.")

table([
    ["Suite", "Count", "Scope"],
    ["Backend - pytest", "50",
     "Route-level tests through FastAPI's TestClient against a temporary SQLite database, plus "
     "pure-function tests. Every test gets its own database file; no shared mutable state."],
    ["Frontend - Vitest", "28",
     "Pure helpers and controlled-component behaviour, mounted with React Testing Library where "
     "interaction matters."],
], widths=[1.5, 0.6, 4.8], size=9)

para(("Testing principles.", "b"), space_before=6)
bullet("Test the logic, not the wiring. ", "Path resolution, split invariants, cache semantics and "
       "queue state transitions are the highest-value targets.")
bullet("No GPU, no heavy imports, no subprocesses. ",
       "Because the machine-learning imports are lazy, a test that never enters a render path "
       "never loads PyTorch. Subprocess boundaries are stubbed.")
bullet("Isolated state. ", "Temporary directories and temporary databases per test; no sleeping.")
bullet("A bug fix ships with a regression test. ",
       "For example, the region-of-interest containment and nearest-window fallback rules each "
       "have a direct unit test.")

h("8.1 Test-case trace table", before=10, after=4)
para("Representative rows from the implemented suites. Where an automated test covers the row, "
     "the last column names it, so the case is executable rather than aspirational. Two queue "
     "rows have no unit test: the compare-and-swap races were verified by construction and "
     "exercised manually with two browser tabs, as recorded in the queue design document.")

table([
    ["TC", "Module", "Scenario", "Expected result", "Automated test / verification"],
    ["TC-01", "Queue / worker", "A job is cancelled in the instant the worker claims it",
     "Exactly one side wins. The conditional update matches zero rows for the loser; the cancel "
     "call returns 409 and the job runs, or the job never starts.",
     "Verified by construction of the compare-and-swap claim; exercised manually with two "
     "browser tabs."],
    ["TC-02", "Queue", "Waiting jobs are dragged into a new order",
     "Positions are rewritten atomically; jobs that changed status mid-reorder are ignored and "
     "unnamed jobs keep their relative order at the back.", "POST /api/queue/reorder"],
    ["TC-03", "Jobs", "A job carries per-job parameters through the queue",
     "The params blob round-trips from enqueue to handler unchanged.",
     "test_enqueue_job_params_roundtrip"],
    ["TC-04", "Splitter", "A single split is created from a manifest",
     "Every data row lands in train.csv; val and test are header-only; the name matches "
     "alltrain_seed_{seed}_{rand8}; metadata records k=1.",
     "test_single_split_all_rows_in_train, test_single_split_name_pattern, "
     "test_single_split_metadata"],
    ["TC-05", "Splitter", "The CSV has no slide-identifier column",
     "The request is rejected before anything is written.",
     "test_single_split_rejects_missing_slide_id, test_single_split_rejects_empty"],
    ["TC-06", "Splitter", "Slide identifiers carry a .tif suffix",
     "The suffix is stripped in the generated CSVs; the source CSV is untouched.",
     "test_single_split_strips_tif"],
    ["TC-07", "Filesystem sandbox", "A path outside the allowed roots is requested",
     "403 Forbidden, before any file is opened.", "test_403_outside_roots"],
    ["TC-08", "Thumbnails", "A slide the imaging library cannot open",
     "422, and the picker falls back to an icon rather than breaking the grid.",
     "test_generation_path_422_on_openslide_failure"],
    ["TC-09", "Thumbnails", "The same slide is requested twice",
     "The cache key is deterministic, so the second request is served from disk.",
     "test_cache_key_deterministic"],
    ["TC-10", "PANTHER training", "A single run is submitted",
     "Exactly one model row is created with no model group; a mismatched dataset or an "
     "unresolvable features directory is rejected.",
     "test_single_run_creates_one_model_no_group, test_single_run_rejects_dataset_mismatch, "
     "test_single_run_rejects_bad_features_dir"],
    ["TC-11", "Per-slide render", "A slide is requested for the comparison page",
     "A cached manifest returns 'ready' with no job; otherwise a render job is enqueued. A "
     "missing feature file or a malformed slide identifier is rejected with 422.",
     "test_render_slide_ready_on_cached_manifest, test_render_slide_enqueues_when_no_manifest, "
     "test_render_slide_422_when_h5_absent, test_render_slide_422_on_bad_slide_id"],
    ["TC-12", "Region of interest", "The user clicks a point on the assignment map",
     "The window containing the point is chosen; a click on a gap falls back to the nearest "
     "window centre; out-of-range coordinates are clamped.",
     "test_render_roi_at_point_containment, test_render_roi_at_point_nearest_fallback, "
     "test_render_roi_at_point_clamps_out_of_range"],
    ["TC-13", "Region of interest", "Selection is attempted before the panel exists",
     "409 Conflict with an explanatory message, not a crash.",
     "test_select_roi_409_when_no_section_a, test_select_roi_409_when_no_compare_manifest, "
     "test_select_roi_409_when_wsi_unavailable"],
    ["TC-14", "Assignment cache", "Several panels render the same (model, slide)",
     "The encoder runs once; the cache evicts past its capacity and never leaks assignments "
     "across models.",
     "test_get_or_compute_calls_compute_only_on_miss, test_lru_eviction_past_cap, "
     "test_per_model_slide_isolation"],
    ["TC-15", "Datasets", "A dataset has only legacy K-fold models",
     "It is invisible to the comparison page, and its slides endpoint returns 404.",
     "test_list_datasets_only_single_models, test_slides_404_when_only_legacy_models"],
    ["TC-16", "Datasets", "The features directory is missing or unreadable",
     "An empty slide list plus an explanatory note - never a 500.",
     "test_slides_missing_features_dir_graceful"],
    ["TC-17", "Frontend - numeric input", "A hyperparameter field is cleared mid-edit",
     "The field can be emptied without the cursor fighting the user; the value clamps to its "
     "range only on blur.",
     "'can be cleared mid-edit...', 'clamps to the range only on blur'"],
    ["TC-18", "Frontend - split scope", "The features directory is re-resolved to the same dataset",
     "The chosen split is preserved - the scope key is the dataset name, not the object identity.",
     "'is STABLE across distinct resolved objects with the same dataset'"],
    ["TC-19", "Frontend - zoom/pan", "A click is mapped to image coordinates under zoom and pan",
     "The transform is inverted correctly and clicks past an edge clamp to the [0,1] range.",
     "'inverts pan and zoom to recover natural coords', 'clamps a click past the left/top edge'"],
    ["TC-20", "Frontend - compare grid", "One to four models are laid out",
     "One to three form a single row; four wrap into a two-by-two grid; overflow is capped.",
     "'wraps 4 models into a 2x2 grid', 'caps at the 2x2 grid for any overflow count'"],
], widths=[0.4, 0.85, 1.5, 2.0, 1.65], size=8)

h("8.2 Failure modes and how the system responds", before=10, after=4)
table([
    ["Failure", "Response"],
    ["A training or rendering subprocess exits non-zero",
     "The traceback is written to the job log, the job is marked failed with the message on the "
     "row, and the model is marked failed. The worker continues with the next job. The user sees "
     "a red pill and a log link."],
    ["A single render inside a multi-render job fails",
     "Each render step is individually guarded, so a partial result still publishes. A missing "
     "panel is simply absent from the artifact manifest, and the interface substitutes a labelled "
     "placeholder."],
    ["A machine-learning dependency is missing from the image",
     "The failure is confined to the render job. Models still train; only the panels are absent. "
     "This is exactly how four missing dependencies were diagnosed and fixed on the GPU box "
     "without the application ever going down."],
    ["A slide has no embedded resolution metadata",
     "The thumbnail renders without a scale bar rather than failing."],
    ["A slide file cannot be found under the recorded directory",
     "The lookup walks the directory subtree by filename stem before giving up, and the affected "
     "panel is skipped with a note in the log."],
    ["A path outside the allowed roots is submitted",
     "403 before any filesystem operation, including when a raw path is pasted into a form."],
    ["The user cancels a job the worker has just claimed",
     "409 with an explanatory message; the job runs to completion. No partial-kill state exists."],
    ["A batch of slides contains some invalid paths",
     "207 Multi-Status; valid paths are queued and invalid ones are reported per path."],
    ["The database is contended",
     "Write-ahead logging plus a five-second busy timeout absorb it."],
], widths=[2.3, 4.6], size=8.5)

doc.add_page_break()

# ==========================================================================
h("9. Monitoring, Troubleshooting, and Integration", before=0)
# ==========================================================================
bullet("User-facing monitoring. ",
       "The queue page is the operational dashboard: what is running now with its elapsed time, "
       "what is waiting and in what order, and the last twenty terminal jobs with their outcome. "
       "Every model, fold and inference additionally carries a status pill and, where relevant, "
       "an inline poller that refreshes the page's data when a job it is watching finishes.")
bullet("Per-job logs. ",
       "Every asynchronous job streams a line-buffered log file, always flushed. The tail is "
       "exposed through the API and rendered in a modal that refreshes while the job is in "
       "flight. This is the primary debugging surface, and it is reachable from every place a "
       "job status is shown, including from a failed model row.")
bullet("Execution records. ",
       "Each training attempt persists its exact command, return code, captured output and "
       "timings in its own table. A failed run remains fully inspectable after the fact.")
bullet("Deliberate degradation. ",
       "The interface never assumes a render exists. Every image path is resolved through a helper "
       "that substitutes a labelled placeholder for a missing or not-yet-rendered artifact, so a "
       "model that is halfway through rendering shows a coherent page rather than broken images.")
bullet("Diagnostics that surface unverified assumptions. ",
       "The slide listing reports how many TRIDENT-written thumbnails were found, so a wrong "
       "assumption about that directory shows up as a visible number rather than as silently "
       "slower thumbnails.")

para(("A symptom-to-cause table for first deployment", "b"),
     " is maintained in the backend operations guide. The recurring causes are: the environment "
     "file not being loaded into the server process; the allow-list naming host paths instead of "
     "container paths; a missing dependency in the shared virtual environment; and an output root "
     "that does not match where the renderers wrote.", space_before=6)

# ==========================================================================
h("10. Documentation Artifacts", before=12)
# ==========================================================================
para("Documentation is versioned with the code, and each document has a single owner subsystem.")
table([
    ["Document", "Contents"],
    ["README.md", "Product overview, screenshots, architecture summary, quick start, credits."],
    ["docs/structure.md", "End-to-end architecture, the domain model, the repository map, and an "
     "explicit statement of what is built versus deferred."],
    ["docs/backend.md", "The working backend reference: stack, every module, the full HTTP API, "
     "the job system, and the deploy-day assumptions to validate."],
    ["docs/frontend.md", "Routing, the typed API client, every page and component, and the "
     "cross-cutting patterns."],
    ["docs/database.md", "The schema, table by table, plus the additive-column convention."],
    ["docs/queue-design.md", "The queue's locked decisions, concurrency model and verification "
     "procedure."],
    ["docs/tests.md", "The testing strategy, fixtures and conventions."],
    ["docs/docker.md", "Host prerequisites, CUDA tag selection, build and run, volume layout, "
     "and a troubleshooting list."],
    ["docs/env-paths.md", "The host-versus-container path model - written because this was the "
     "largest source of deployment confusion."],
    ["docs/test-run.md", "A recorded end-to-end run on the GPU machine."],
    ["This document", "The detailed design, revision 2."],
], widths=[1.6, 5.3], size=8.5)

# ==========================================================================
h("11. Timeline: plan versus actual", before=12)
# ==========================================================================
para("The rev. 1 plan ran March to June in five phases. The work actually ran from 10 April to "
     "23 July 2026 across 39 commits. The phase structure held; the content of two phases changed.")

table([
    ["Phase", "Planned", "Actual"],
    ["1 - Infrastructure",
     "Weeks 1-3. Run PANTHER on dummy feature vectors on the hospital GPU server.",
     "10-18 April. Built as a mock-worker prototype with an annotation and comparison user "
     "interface, deliberately without touching a GPU, to settle the interaction model first."],
    ["2 - Core backend",
     "Weeks 4-7. Decoupled asynchronous polling workers; connect to the hospital database; stream "
     "patches to the client.",
     "19 May - 24 May. The project was restructured around the real TRIDENT/PANTHER pipeline, the "
     "frontend migrated to TypeScript, and the mock worker removed. K-fold training, the "
     "asynchronous worker and the job table landed here. The hospital-database integration was "
     "replaced by an application-owned SQLite schema."],
    ["3 - Visualization",
     "Weeks 8-11. Cluster visualization, heatmaps, and a pop-out feature.",
     "24 May - 30 June. Delivered as four analysis panels reproducing the paper's figure layout, "
     "plus a hand-rolled zoom/pan viewer. This phase absorbed the most iteration: the region-of-"
     "interest ranking, the true patch pitch and the tissue framing were each corrected against "
     "real slides."],
    ["4 - Researcher tools",
     "Weeks 12-14. Semantic labelling and annotations.",
     "8 June - 27 June. Labelling and notes, the shared queue page with cancel and drag-reorder, "
     "the design system, model deletion, and CSV validation."],
    ["5 - Deployment, testing, submission",
     "Weeks 15-16. Stress testing, debugging, submission.",
     "24 June - 23 July. Containerization and the first real GPU deployment, which produced the "
     "corrections recorded in section 0; then the shift from K-fold to single-model training, the "
     "Model Comparison page, and the automated test suites."],
], widths=[1.15, 2.2, 3.55], size=8.5)

para(("What the plan got wrong, and why it is worth recording. ", "b"),
     "The estimate assumed that integrating with an existing pipeline is mostly plumbing. In "
     "practice, roughly a third of the effort went into a category the plan did not name: "
     "reconciling assumptions about what upstream code expects on disk. PANTHER reads its splits "
     "from a directory one level deeper than documented, asserts on the name of the features "
     "directory, and needs a CUDA-specific clustering library; TRIDENT names its output "
     "differently from what PANTHER accepts; the feature width is a property of the encoder rather "
     "than a hyperparameter the user should be asked for. None of these are visible from reading "
     "the papers. The mitigation that worked was structural rather than schedule-based: because "
     "every one of these failures surfaced as a failed job with a readable log while the "
     "application stayed up, each was diagnosed in a single iteration.")

# ==========================================================================
h("12. Known Limitations and Future Work", before=12)
# ==========================================================================
para("These are deliberate, documented boundaries rather than defects.")
bullet("No job cancellation once running. ",
       "A single thread cannot interrupt a subprocess cleanly without leaving partial output. "
       "Waiting jobs can be cancelled; running ones are left to finish. Removing this limitation "
       "means replacing the single thread with a real queue and worker pool.")
bullet("Single user, single GPU. ",
       "One worker thread serialises all GPU work. The queue is shared and any user may cancel or "
       "reorder any job, but there is no authentication and no per-user isolation.")
bullet("No authentication or access control. ",
       "Acceptable only because the system is reached over a controlled internal network. "
       "This is the first thing to add before any wider deployment.")
bullet("Output directories grow without bound. ",
       "Rendered figures, job logs and inference outputs are never pruned automatically. A cleanup "
       "endpoint is sketched but not implemented.")
bullet("Slides cannot be uploaded through the browser. ",
       "Files must already exist on the server within an allowed root.")
bullet("Feature extraction blocks its request. ",
       "The initial TRIDENT run is synchronous. It predates the job system and was left that way "
       "because it happens once per dataset; moving it onto the queue is straightforward.")
bullet("No resolution override. ",
       "Slides without embedded physical resolution render without a scale bar; there is no way to "
       "supply the value manually.")
bullet("Standalone models cannot be deleted from the interface. ",
       "The cascading delete is keyed on a model-group row, which single runs do not create. "
       "Legacy K-fold groups can be deleted; a standalone model must be removed by hand. This "
       "is the one place where the shift away from K-fold left a gap.")
bullet("No export. ",
       "Prototype labels, notes and inference results can be read in the interface but not "
       "exported to a file.")
bullet("K-fold is legacy. ",
       "Existing groups remain viewable but no new K-fold runs can be created. Reviving it as a "
       "first-class flow would be a small change to the training route.")
bullet("One assumption remains unverified. ",
       "The exact subdirectory TRIDENT writes slide thumbnails to. The system probes three "
       "extensions, falls back to generating thumbnails itself, and reports a diagnostic count so "
       "the mismatch is visible.")

para(("Closing note. ", "b"),
     "The architecture that this document describes was chosen for one property above all others: "
     "that a failure anywhere in the machine-learning pipeline degrades into a red status and a "
     "readable log, rather than into an unresponsive application. Every significant problem "
     "encountered during deployment was diagnosed through that mechanism. If one design decision "
     "from this project is worth carrying forward, it is the decision to let a database table, "
     "rather than a function call, be the boundary between the interface and the model.")

doc.save(OUT)
print("wrote", OUT)
